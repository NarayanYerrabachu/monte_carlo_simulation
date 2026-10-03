"""PDF and Excel reports for a finished job (fleet simulation, scenario simulation or loop test).

Built only from the job's stored result — the same JSON the live viewer and the
report page show — so every format says the same thing. Fleet and scenario
reports also carry the TDA Mapper graph the job was sent (3D picture + the
labelled groups), i.e. they are "Monte Carlo with TDA" reports.
"""
from __future__ import annotations

import io
from datetime import UTC, datetime
from functools import reduce
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.ticker import FuncFormatter, MaxNLocator
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from mc_service.fleet import (
    charts as fleet_charts,
)
from mc_service.fleet.pdf import S, _footer, _img, _section, _table_style

BLUE, ORANGE, GREEN, RED, INK = fleet_charts.BLUE, fleet_charts.ORANGE, fleet_charts.AQUA, "#d03b3b", fleet_charts.INK
HEAD_FILL, HEAD_FONT = PatternFill("solid", fgColor="1C5CAB"), Font(bold=True, color="FFFFFF")


def pct(v: float | None, d: int = 1) -> str:
    return "–" if v is None else f"{v * 100:.{d}f}%"


def num(v: float | None, d: int = 0) -> str:
    return "–" if v is None else f"{v:,.{d}f}"


def eur(v: float | None) -> str:
    return "–" if v is None else f"€{v:,.0f}"


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _bars(ax, h: dict, color: str, lines: list[tuple[float, str, str]], xlabel: str) -> None:
    edges, counts = h["edges"], h["counts"]
    total = sum(counts) or 1
    width = (edges[1] - edges[0]) * 0.9 if len(edges) > 1 else 1
    ax.bar(edges[:-1], [c / total for c in counts], width=width, align="edge", color=color)
    top = max(counts) / total if counts else 1
    for i, (x, label, col) in enumerate(lines):
        ax.axvline(x, color=col, lw=1.3, ls="--")
        ax.annotate(label, (x, top * (1.25 - 0.12 * i)), xytext=(4, 0), textcoords="offset points",
                    fontsize=7.5, color=col, fontweight="bold", va="center")
    ax.set_ylim(0, top * 1.35)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=5))               # large values (costs) must not overlap
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}" if abs(v) >= 1000 else f"{v:g}"))
    ax.set_xlabel(xlabel)
    ax.set_ylabel("probability")
    ax.grid(axis="x", visible=False)


# ── the report page's layout, shared by PDF and Excel ───────────────────────
# A report is described once as "blocks" — the same boxes the HTML report shows:
#   inputs   [(label, value)]             1. Uncertain inputs
#   engine   (count, unit, [notes])       Monte Carlo engine
#   outputs  [(kind, value, text)]        2. Probabilistic outputs (kind: good | bad | info)
#   tiles    [(label, value, sub, tone)]  KPI tiles (tone: good | blue | warm | bad)
TONE_HEX = {"good": "#15803d", "blue": "#2563eb", "warm": "#d97706", "bad": "#b91c1c", "info": "#2563eb"}
PANEL, BORDER = rl_colors.HexColor("#f1f3f6"), rl_colors.HexColor("#dfe3e0")
_P = {
    "box_h": ParagraphStyle("box_h", parent=S["cell"], fontName="DejaVu-Bold", fontSize=9, leading=12),
    "kv_k": ParagraphStyle("kv_k", parent=S["small"], fontSize=6.5, leading=8.5),
    "kv_v": ParagraphStyle("kv_v", parent=S["cell"], fontName="DejaVu-Bold", fontSize=8, leading=10.5),
    "engine": ParagraphStyle("engine", parent=S["cell"], alignment=1, fontSize=7.5, leading=10),
    "engine_n": ParagraphStyle("engine_n", parent=S["cell"], alignment=1, fontName="DejaVu-Bold", fontSize=14,
                               leading=18, textColor=rl_colors.HexColor(TONE_HEX["warm"])),
    "tile_k": ParagraphStyle("tile_k", parent=S["small"], fontSize=6.5, leading=8.5),
    "tile_s": ParagraphStyle("tile_s", parent=S["small"], fontSize=7, leading=9),
}


def _flow(blocks: dict) -> Table:
    """1. Uncertain inputs → Monte Carlo engine → 2. Probabilistic outputs, as on the report page."""
    pairs = blocks["inputs"]
    cells = [[Paragraph(k.upper(), _P["kv_k"]), Paragraph(v, _P["kv_v"])] for k, v in pairs]
    grid = Table([[cells[i], cells[i + 1] if i + 1 < len(cells) else ""] for i in range(0, len(cells), 2)],
                 colWidths=[2.65 * cm, 2.65 * cm])
    grid.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), PANEL), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("LINEBELOW", (0, 0), (-1, -1), 2, rl_colors.white),
                              ("LINEAFTER", (0, 0), (0, -1), 2, rl_colors.white),
                              ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                              ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
    count, unit, notes = blocks["engine"]
    engine = [Paragraph("MONTE CARLO ENGINE", S["eyebrow"]), Spacer(1, 4), Paragraph(count, _P["engine_n"]),
              Paragraph(f"simulated {unit}", _P["engine"]), Spacer(1, 4), *[Paragraph(n, _P["engine"]) for n in notes]]
    marks = {"good": "✓", "bad": "!", "info": "•"}
    outs = [Paragraph(f'<font color="{TONE_HEX[kind]}" name="DejaVu-Bold">{marks[kind]}</font>&nbsp; <font name="DejaVu-Bold">{value}</font> {text}', S["cell"])
            for kind, value, text in blocks["outputs"]]
    out_table = Table([[o] for o in outs], colWidths=[6.9 * cm])
    out_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), PANEL), ("LINEBELOW", (0, 0), (-1, -1), 2, rl_colors.white),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
    head = lambda n, text, sub="": Paragraph(
        f'<font color="{TONE_HEX["warm"]}">{n}.</font> {text} <font color="#5f6b66" size="7">{sub}</font>', _P["box_h"])
    arrow = Paragraph("→", ParagraphStyle("arrow", parent=S["cell"], alignment=1, fontSize=13, textColor=rl_colors.HexColor("#5f6b66")))
    t = Table([[[head(1, "Uncertain inputs", blocks.get("inputs_sub", "observed in the data")), Spacer(1, 4), grid], arrow,
                engine, arrow, [head(2, "Probabilistic outputs"), Spacer(1, 4), out_table]]],
              colWidths=[5.6 * cm, 0.5 * cm, 3.1 * cm, 0.5 * cm, 7.3 * cm])
    t.setStyle(TableStyle([("BOX", (0, 0), (0, 0), 0.6, BORDER), ("BOX", (2, 0), (2, 0), 0.6, BORDER),
                           ("BOX", (4, 0), (4, 0), 0.6, BORDER), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("VALIGN", (1, 0), (1, 0), "MIDDLE"), ("VALIGN", (3, 0), (3, 0), "MIDDLE"),
                           ("VALIGN", (2, 0), (2, 0), "MIDDLE"),
                           ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                           ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return t


def _tiles(blocks: dict) -> Table:
    """The KPI tiles row under the flow."""
    tiles = blocks["tiles"]
    cells = [[Paragraph(label.upper(), _P["tile_k"]),
              Paragraph(value, ParagraphStyle("tile_v", parent=S["big"], fontSize=16, leading=20,
                                              textColor=rl_colors.HexColor(TONE_HEX.get(tone, TONE_HEX["blue"])))),
              Paragraph(sub_text, _P["tile_s"])] for label, value, sub_text, tone in tiles]
    gap, n = 0.25 * cm, len(tiles)
    width = (17 * cm - gap * (n - 1)) / n
    row, widths = [], []
    for i, cell in enumerate(cells):
        row.append(cell)
        widths.append(width)
        if i < n - 1:
            row.append("")
            widths.append(gap)
    t = Table([row], colWidths=widths)
    style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
             ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
    style += [("BOX", (2 * i, 0), (2 * i, 0), 0.6, BORDER) for i in range(n)]
    t.setStyle(TableStyle(style))
    return t


def _report_sheet(wb: Workbook, title: str, subtitle: str, blocks: dict) -> None:
    """First Excel sheet: the report page top to bottom — KPI tiles, uncertain inputs, engine, outputs."""
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = subtitle
    row = 4
    ws.cell(row=row, column=1, value="Key figures").font = Font(bold=True, size=12)
    _head(ws, row + 1, ["KPI", "Value", "Note"])
    for i, (label, value, sub_text, tone) in enumerate(blocks["tiles"], row + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value).font = Font(bold=True, size=13, color=TONE_HEX.get(tone, TONE_HEX["blue"])[1:])
        ws.cell(row=i, column=3, value=sub_text)
    row += len(blocks["tiles"]) + 3
    ws.cell(row=row, column=1, value=f"1. Uncertain inputs ({blocks.get('inputs_sub', 'observed in the data')})").font = Font(bold=True, size=12)
    _head(ws, row + 1, ["Input", "Observed"])
    for i, (label, value) in enumerate(blocks["inputs"], row + 2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
    row += len(blocks["inputs"]) + 3
    count, unit, notes = blocks["engine"]
    ws.cell(row=row, column=1, value="Monte Carlo engine").font = Font(bold=True, size=12)
    ws.cell(row=row + 1, column=1, value=f"simulated {unit}")
    ws.cell(row=row + 1, column=2, value=count).font = Font(bold=True, size=13, color=TONE_HEX["warm"][1:])
    for i, note in enumerate(notes, row + 2):
        ws.cell(row=i, column=1, value=note)
    row += len(notes) + 3
    ws.cell(row=row, column=1, value="2. Probabilistic outputs").font = Font(bold=True, size=12)
    _head(ws, row + 1, ["Value", "Meaning"])
    for i, (kind, value, text) in enumerate(blocks["outputs"], row + 2):
        ws.cell(row=i, column=1, value=value).font = Font(bold=True, color=TONE_HEX[kind][1:])
        ws.cell(row=i, column=2, value=text)
    _widths(ws, [46, 34, 70])


MAX_REGIME_ROWS = 12          # regimes listed in the PDF; the Excel sheet "TDA regimes" has all of them


def _regime_note(n: int, order: str) -> str | None:
    if n <= MAX_REGIME_ROWS:
        return None
    return (f"The {MAX_REGIME_ROWS} most relevant of {n} regimes, {order}; the Excel report lists all of them "
            "(sheet “TDA regimes”).")


# ── fleet ───────────────────────────────────────────────────────────────────
def _fleet_blocks(f: dict) -> dict:
    k, c, inp = f["summary"]["kpi"], f["config"], f["summary"].get("inputs") or {}
    r = lambda q, fmt: "–" if not q else f"{fmt(q['p5'])} – {fmt(q['p95'])}"
    return {
        "inputs": [("Vehicles in data", f"{num(inp.get('vehicles'))} · {num(inp.get('days'))} days"),
                   ("Daily demand (P5–P95)", r(inp.get("daily_demand"), num)),
                   ("Driver availability", r(inp.get("driver_availability"), lambda v: pct(v, 0))),
                   ("Breakdowns per day", r(inp.get("daily_breakdowns"), num)),
                   ("Fuel consumption", f"{r(inp.get('fuel_l_per_100km'), lambda v: num(v, 1))} L/100 km"),
                   ("Route duration", f"{r(inp.get('route_duration_h'), lambda v: num(v, 1))} h"),
                   ("Fuel price", f"{r(inp.get('fuel_price_eur_l'), lambda v: num(v, 2))} €/L"),
                   ("Maintenance per breakdown", r(inp.get("maintenance_cost_per_breakdown"), eur))],
        "engine": (num(f["n_completed"]), "operating days",
                   [(f"{num(c['n_pool'])} of {num(c['n_records'])} records · {c['n_days']} historical days · fleet of "
                     f"{num(c['fleet_size'])}"), "Each scenario draws a real day and resamples its vehicle records"]),
        "outputs": [
            ("good" if k["p_meet_sla"] >= 0.9 else "bad", pct(k["p_meet_sla"]),
             f"probability a day meets the on-time SLA (≥ {pct(c['sla_on_time'], 0)} of deliveries on time)"),
            ("bad" if k["p_miss_sla"] > 0.1 else "good", pct(k["p_miss_sla"]), "probability of missing the delivery target"),
            ("info", f"{num(k['vehicles_required_mean'])} vehicles",
             f"expected to meet demand (95th percentile: {num(k['vehicles_required_p95'])}; fleet: {num(k['fleet_size'])})"),
            ("info", f"{num(k['fuel_l_mean'])} L", f"expected daily fuel consumption (P95 {num(k['fuel_l_p95'])} L)"),
            ("info", pct(k["p_breakdowns_over_alert"]),
             f"probability of more than {k['breakdown_alert']} vehicle breakdowns in a day"),
            ("info", eur(k["maint_cost_mean"]), f"expected maintenance cost per day (P95 {eur(k['maint_cost_p95'])})")],
        "tiles": [("Fleet availability", pct(k["fleet_availability_mean"]), "vehicles with driver and no breakdown", "good"),
                  (f"Delivered within {c['delivery_target_h']:g} h", pct(k["p_delivery_within_target"]),
                   "of all planned deliveries (missed ones count as late)", "blue"),
                  ("Expected daily fuel cost", eur(k["fuel_cost_mean"]), f"P95 {eur(k['fuel_cost_p95'])}", "warm"),
                  ("Breakdown risk", pct(k["p_breakdowns_over_alert"]), f"P(> {k['breakdown_alert']} breakdowns per day)", "bad")],
    }


def _fleet_charts(f: dict) -> dict[str, bytes]:
    k, c, s = f["summary"]["kpi"], f["config"], f["summary"]
    out = {}
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.7))
    dt = s["delivery_time"]
    if dt["share"]:
        mids = [(a + b) / 2 for a, b in zip(dt["edges"], dt["edges"][1:], strict=False)]
        axes[0].bar(mids, dt["share"], width=(dt["edges"][1] - dt["edges"][0]) * 0.9,
                    color=[BLUE if m <= c["delivery_target_h"] else "#b9c0cb" for m in mids])
        axes[0].axvline(c["delivery_target_h"], color=RED, ls="--", lw=1.3)
        axes[0].set_title(f"Delivery time · P(within {c['delivery_target_h']} h | delivered) = {pct(dt['p_within_target'])}",
                          fontsize=8, color=INK, loc="left")
        axes[0].set_xlabel("delivery time (hours)")
        axes[0].set_ylabel("probability")
    o = s["outcome"]
    axes[1].pie([o["on_time"], o["delayed"]], labels=["On time", "Delayed"], colors=[GREEN, RED],
                autopct="%1.1f%%", startangle=90, counterclock=False, wedgeprops={"width": 0.45},
                textprops={"fontsize": 8})
    axes[1].set_title("Delivery outcome", fontsize=8, color=INK)
    fig.tight_layout()
    out["delivery"] = _png(fig)

    fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.7))
    _bars(axes[0], s["hist"]["vehicles_required"], BLUE,
          [(k["vehicles_required_mean"], f"expected {num(k['vehicles_required_mean'])}", ORANGE),
           (k["vehicles_required_p95"], f"P95 {num(k['vehicles_required_p95'])}", RED),
           (k["fleet_size"], f"fleet {num(k['fleet_size'])}", GREEN)], "vehicles required")
    _bars(axes[1], s["hist"]["maint_cost"], GREEN,
          [(k["maint_cost_mean"], f"expected {eur(k['maint_cost_mean'])}", ORANGE),
           (k["maint_cost_p95"], f"P95 {eur(k['maint_cost_p95'])}", RED)], "maintenance cost per day (€)")
    fig.tight_layout()
    out["dists"] = _png(fig)

    cv = s["convergence"]
    fig, ax = plt.subplots(figsize=(7.2, 2.2))
    ax.fill_between(cv["n"], cv["lo"], cv["hi"], color=ORANGE, alpha=0.2, label="95% confidence interval")
    ax.plot(cv["n"], cv["p"], color=ORANGE, lw=2, label="P(meet SLA)")
    ax.set_xscale("log")
    ax.set_xlabel("simulated days (log scale)")
    ax.legend(frameon=False, fontsize=8, loc="lower right")
    out["convergence"] = _png(fig)
    return out


def _fleet_story(f: dict, job: dict) -> list:
    c, s = f["config"], f["summary"]
    png = _fleet_charts(f)
    ctx = s.get("context") or {}
    story: list = [
        Paragraph("FLEET MANAGEMENT · MONTE CARLO ON TDA / ML RESULTS", S["eyebrow"]),
        Paragraph("Monte Carlo Simulation in Fleet Management", S["title"]),
        Spacer(1, 4),
        Paragraph(f"Dataset {job['dataset_id']} · job {job['job_id']} · {num(f['n_completed'])} simulated operating days · "
                  f"generated {datetime.now(UTC):%d %b %Y}", S["small"]),
    ]
    blocks = _fleet_blocks(f)
    story += [Spacer(1, 10), _flow(blocks), Spacer(1, 8), _tiles(blocks), Spacer(1, 8),
              Paragraph(f"Each simulated day draws a real historical day and resamples {num(c['fleet_size'])} vehicle records "
                        f"from it ({num(c['n_pool'])} of {num(c['n_records'])} records in the pool). Demand stays at the "
                        "historical level; a vehicle's observed on-time deliveries are its capacity.", S["small"])]

    story += [KeepTogether(_section("Charts", "Distributions") + [_img(png["delivery"])]), _img(png["dists"])]

    reg = [["Regime (TDA cluster)", "Share", "On time", "Breakdown", "Availability", "Late share"]] + [
        [Paragraph(g["label"], S["cell"]), pct(g["share"]), pct(g["on_time_share"]), pct(g["breakdown_rate"]),
         pct(g["availability"]), pct(g["late_deliveries_share"])]
        for g in sorted(f["results"], key=lambda g: -g["late_deliveries_share"])[:MAX_REGIME_ROWS]]
    tr = Table(reg, colWidths=[6.2 * cm, 2 * cm, 2.2 * cm, 2.2 * cm, 2.2 * cm, 2.2 * cm], repeatRows=1)
    tr.setStyle(_table_style())
    story += [KeepTogether(_section("TDA", "Where the risk comes from — TDA regimes", _regime_note(len(f["results"]),
                                    "by their share of the late deliveries")) + [tr])]

    notes = []
    if ctx:
        notes.append(f"TDA / ML: {num(ctx.get('n_clusters'))} clusters, {num(ctx.get('n_noise'))} noise records, "
                     f"{num(ctx.get('n_anomalies_high'))} records scored HIGH by the anomaly model.")
        for title, items in (("Relationships", ctx.get("relationships")), ("Patterns", ctx.get("patterns"))):
            if items:
                notes.append(f"{title}: " + "; ".join(items))
    a = s["anomalies"]
    notes.append("Anomaly exclusion off: all records used." if a["threshold"] is None else
                 f"{num(a['excluded'])} operating records with anomaly score ≥ {a['threshold']} excluded; "
                 f"{num(a.get('kept_events'))} flagged breakdown / absence records kept.")
    story += _mapper_story(job.get("graph"))
    story += [KeepTogether(_section("Context", "TDA / ML findings and data quality")
                           + [Paragraph(n, S["body"]) for n in notes])]
    story += [KeepTogether(_section("Quality", "Are the scenarios enough?") + [_img(png["convergence"])])]
    return story


def _fleet_xlsx(f: dict, job: dict, wb: Workbook) -> None:
    k, c, s = f["summary"]["kpi"], f["config"], f["summary"]
    subtitle = f"Dataset {job['dataset_id']} · job {job['job_id']} · {f['n_completed']:,} simulated days"
    _report_sheet(wb, "Monte Carlo Simulation in Fleet Management", subtitle, _fleet_blocks(f))
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Monte Carlo Simulation in Fleet Management"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = subtitle
    _head(ws, 4, ["Metric", "Value"])
    rows = [("P(day meets SLA)", k["p_meet_sla"], "0.0%"), ("P(miss target)", k["p_miss_sla"], "0.0%"),
            ("SLA (on-time share)", c["sla_on_time"], "0%"), ("Delivery target (h)", c["delivery_target_h"], "0.0"),
            ("Delivered within target (of planned)", k["p_delivery_within_target"], "0.0%"),
            ("Fleet size", k["fleet_size"], "#,##0"), ("Vehicles required (mean)", k["vehicles_required_mean"], "#,##0.0"),
            ("Vehicles required (P95)", k["vehicles_required_p95"], "#,##0"),
            ("Fleet availability", k["fleet_availability_mean"], "0.0%"),
            ("Fuel per day (L, mean)", k["fuel_l_mean"], "#,##0"), ("Fuel per day (L, P95)", k["fuel_l_p95"], "#,##0"),
            ("Fuel cost per day (mean)", k["fuel_cost_mean"], "€#,##0"),
            ("Breakdowns per day (mean)", k["breakdowns_mean"], "#,##0.0"),
            (f"P(> {k['breakdown_alert']} breakdowns)", k["p_breakdowns_over_alert"], "0.0%"),
            ("Maintenance cost per day (mean)", k["maint_cost_mean"], "€#,##0"),
            ("Maintenance cost per day (P95)", k["maint_cost_p95"], "€#,##0")]
    for i, (name, value, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=value).number_format = fmt
    _widths(ws, [44, 16])

    ws = wb.create_sheet("Inputs (observed)")
    _head(ws, 1, ["Input", "P5", "Median", "P95"])
    for i, (key, q) in enumerate(((k2, v) for k2, v in (s.get("inputs") or {}).items() if isinstance(v, dict)), 2):
        ws.cell(row=i, column=1, value=key.replace("_", " "))
        for j, part in enumerate(("p5", "median", "p95"), 2):
            ws.cell(row=i, column=j, value=q[part]).number_format = "#,##0.00"
    _widths(ws, [34, 14, 14, 14])

    ws = wb.create_sheet("TDA regimes")
    cols = ["regime", "label", "records", "share", "on_time_share", "within_target_share", "breakdown_rate",
            "availability", "fuel_l_mean", "maint_cost_mean", "late_deliveries_share"]
    _head(ws, 1, [col.replace("_", " ") for col in cols])
    for i, g in enumerate(f["results"], 2):
        for j, col in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j, value=g[col])
            if col.endswith(("share", "rate")) or col == "availability":
                cell.number_format = "0.0%"
    _widths(ws, [8, 36] + [14] * (len(cols) - 2))

    ws = wb.create_sheet("Distributions")
    col = 1
    for key, title in (("vehicles_required", "Vehicles required"), ("maint_cost", "Maintenance cost per day"),
                       ("fuel_l", "Fuel per day (L)"), ("on_time_share", "On-time share per day")):
        h = s["hist"][key]
        total = sum(h["counts"]) or 1
        ws.cell(row=1, column=col, value=title).font = Font(bold=True)
        _head(ws, 2, ["From", "To", "Share"], col)
        for i, (a, b, cnt) in enumerate(zip(h["edges"], h["edges"][1:], h["counts"], strict=False), 3):
            ws.cell(row=i, column=col, value=a)
            ws.cell(row=i, column=col + 1, value=b)
            ws.cell(row=i, column=col + 2, value=cnt / total).number_format = "0.0%"
        chart = BarChart()
        chart.title, chart.legend, chart.gapWidth = title, None, 20
        chart.add_data(Reference(ws, min_col=col + 2, min_row=2, max_row=2 + len(h["counts"])), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=col, min_row=3, max_row=2 + len(h["counts"])))
        chart.height, chart.width = 7, 13
        ws.add_chart(chart, f"{get_column_letter(col)}{len(h['counts']) + 5}")
        col += 4

    ws = wb.create_sheet("Delivery time")
    dt = s["delivery_time"]
    _head(ws, 1, ["From (h)", "To (h)", "Share of deliveries"])
    for i, (a, b, v) in enumerate(zip(dt["edges"], dt["edges"][1:], dt["share"], strict=False), 2):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
        ws.cell(row=i, column=3, value=v).number_format = "0.0%"
    _widths(ws, [10, 10, 20])

    _mapper_sheet(wb, job.get("graph"))
    _convergence_sheet(wb, s["convergence"], "P(meet SLA)")
    _context_sheet(wb, s.get("context"))


# ── TDA Mapper (fleet and scenario reports) ─────────────────────────────────
MAPPER_CMAP = LinearSegmentedColormap.from_list("mapper", ["#5E9CA6", "#9B7FD4", "#E05252"])


def _th(labels: list[str]) -> list:
    """Table header cells that wrap instead of running into each other."""
    return [Paragraph(f"<b>{label}</b>", S["small"]) for label in labels]


def _adaptive(v: float | None) -> str:
    """12,345 · 123.4 · 12.35 · 0.123"""
    if v is None:
        return "–"
    a = abs(v)
    return f"{v:,.{0 if a >= 1000 else 1 if a >= 100 else 2 if a >= 1 else 3}f}"


def _fv(v: float | None, fmt: str, unit: str | None = None) -> str:
    """A value in the format a job asked for ("pct" | "int" | "num"), with its unit."""
    if v is None:
        return "–"
    text = pct(v) if fmt == "pct" else num(v) if fmt == "int" else _adaptive(v)
    if not unit or fmt == "pct":
        return text
    return f"{unit}{text}" if unit in ("€", "$", "£") else f"{text} {unit}"


def _graph_view(g: dict) -> dict[str, Any]:
    """How the job's Mapper groups are coloured, flagged and described (same rules as
    frontend/js/mapper_view.js): scenario jobs send ``graph.view``; fleet jobs use the
    on-time share of the group's deliveries."""
    v = g.get("view")
    if not v:
        return {"color": lambda nd: 0.0 if nd.get("on_time") is None else 1 - nd["on_time"],
                "color_title": "late deliveries", "color_max": 1.0, "risky": lambda nd: nd.get("on_time") is not None and nd["on_time"] < 0.7,
                "columns": [("On time", lambda nd: pct(nd.get("on_time"))),
                            ("Breakdowns", lambda nd: pct(nd.get("breakdown_rate")))],
                "risk_note": "⚠ below 70 % on time"}
    get = lambda nd, path: reduce(lambda o, k: None if o is None else o.get(k), path.split("."), nd)
    return {"color": lambda nd: get(nd, v["color_key"]) or 0.0, "color_title": v["color_title"],
            "color_max": v.get("color_max") or 1.0,
            "risky": lambda nd: (get(nd, v["color_key"]) or 0.0) >= v["risk_above"],
            "columns": [(c["label"], lambda nd, c=c: _fv(get(nd, c["key"]), c["fmt"])) for c in v["columns"]],
            "risk_note": f"⚠ {v['color_title']} ≥ {v['risk_above']}"}


def _labelled(g: dict, gv: dict, limit: int = 12) -> list[dict]:
    """Each distinct label once (on its largest group): the 8 largest plus the groups at risk."""
    seen, shown = set(), []
    for nd in sorted(g["nodes"], key=lambda nd: -nd["size"]):
        name = nd.get("label") or f"group {nd['id']}"
        if len(shown) >= limit:
            break
        if name in seen or nd["size"] < 20:
            continue
        if len(shown) < 8 or gv["risky"](nd):
            seen.add(name)
            shown.append(nd)
    return shown


def _mapper_png(g: dict, gv: dict, shown: list[dict]) -> bytes:
    dark = "#0F1117"
    fig = plt.figure(figsize=(7.2, 4.4), facecolor=dark)
    ax = fig.add_subplot(111, projection="3d", facecolor=dark)
    by_id = {nd["id"]: nd for nd in g["nodes"]}
    for a, b in g["edges"]:
        p, q = by_id.get(a), by_id.get(b)
        if p and q:
            ax.plot([p["x"], q["x"]], [p["y"], q["y"]], [p["z"], q["z"]], color="#AAB6CC", lw=0.6, alpha=0.6)
    sc = ax.scatter([nd["x"] for nd in g["nodes"]], [nd["y"] for nd in g["nodes"]], [nd["z"] for nd in g["nodes"]],
                    s=[max(5, min(28, nd["size"] ** 0.5 * 2.4)) ** 2 * 0.35 for nd in g["nodes"]],
                    c=[gv["color"](nd) for nd in g["nodes"]], cmap=MAPPER_CMAP, vmin=0, vmax=gv["color_max"], alpha=0.9,
                    edgecolors=dark, linewidths=0.4, depthshade=False)
    for i, nd in enumerate(shown, 1):
        ax.text(nd["x"], nd["y"], nd["z"], f"{'⚠' if gv['risky'](nd) else ''}{i}", color="#E6E9EE", fontsize=7,
                ha="center", va="bottom")
    ax.set_axis_off()
    ax.view_init(elev=22, azim=45)
    cb = fig.colorbar(sc, ax=ax, shrink=0.55, pad=0.02)
    cb.set_label(gv["color_title"], color="#B6BDCA", fontsize=7)
    cb.ax.tick_params(colors="#B6BDCA", labelsize=6)
    cb.outline.set_visible(False)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor=dark, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


GALAXY_BG, GALAXY_DISK = "#000008", 0.4


def _galaxy_rgb(t: float) -> tuple[float, float, float]:
    """Gold core → teal arm → supernova red (same scale as frontend/js/mapper_view.js)."""
    mix = lambda a, b, u: tuple((x + (y - x) * u) / 255 for x, y in zip(a, b, strict=True))
    if t < 0.4:
        return mix((255, 200, 50), (0, 220, 200), t / 0.4)
    if t < 0.72:
        return mix((0, 220, 200), (210, 70, 50), (t - 0.4) / 0.32)
    return mix((210, 70, 50), (255, 30, 20), min(1.0, (t - 0.72) / 0.28))


def _galaxy_png(g: dict, gv: dict, shown: list[dict]) -> bytes:
    """The Galaxy view of the report page as a picture: the Mapper groups flattened into a disk,
    with a starfield, nebula gas (disk haze, a cloud per group, streaks along the connections)
    and a glow per group. Colour: gold (low risk) → teal → red (high).

    Drawn as a projected 2D picture (the disk seen from 35° above, turned 45°) so the layers
    stack in a fixed order — stars, gas, connections, glow, groups, numbers. matplotlib's 3D
    axes sort whole collections by depth, which puts the gas on top of the groups."""
    rng = np.random.default_rng(77)
    nodes = g["nodes"]
    span = max(1e-6, max(float(np.hypot(nd["x"], nd["y"])) for nd in nodes))
    tint = {nd["id"]: _galaxy_rgb(min(1.0, gv["color"](nd) / (gv["color_max"] or 1.0))) for nd in nodes}
    px = {nd["id"]: max(5.0, min(28.0, nd["size"] ** 0.5 * 2.4)) for nd in nodes}
    by_id = {nd["id"]: nd for nd in nodes}
    turn, tilt = np.radians(45), np.radians(35)

    def project(x, y, z):
        """Data coordinates (z already flattened) → picture coordinates."""
        x, y, z = np.asarray(x, dtype=float), np.asarray(y, dtype=float), np.asarray(z, dtype=float)
        xr, yr = x * np.cos(turn) - y * np.sin(turn), x * np.sin(turn) + y * np.cos(turn)
        return xr, yr * np.sin(tilt) + z * np.cos(tilt)

    scale = 0.5                                   # viewer marker pixels → points on this figure
    fig, ax = plt.subplots(figsize=(7.2, 4.2), facecolor=GALAXY_BG)
    ax.set_facecolor(GALAXY_BG)

    # starfield: a flattened shell around the galaxy
    n_st = 420
    rad, th, ph = span * rng.uniform(1.1, 1.9, n_st), rng.uniform(0, 2 * np.pi, n_st), rng.uniform(0, np.pi, n_st)
    sx, sy = project(rad * np.sin(ph) * np.cos(th), rad * np.sin(ph) * np.sin(th), rad * np.cos(ph) * 0.18)
    ax.scatter(sx, sy, s=rng.uniform(0.1, 1.4, n_st), c="white", alpha=0.8, linewidths=0, zorder=1)

    # gas: disk haze, a cloud in each group's colour, streaks along the connections
    gx, gy, gz, gc, gs = [], [], [], [], []
    n_haze = 360
    rad, th = span * 0.6 * np.abs(rng.normal(size=n_haze)), rng.uniform(0, 2 * np.pi, n_haze)
    gx += (rad * np.cos(th)).tolist()
    gy += (rad * np.sin(th)).tolist()
    gz += rng.normal(0, 0.05, n_haze).tolist()
    gc += [(rng.uniform(60, 100) / 255, rng.uniform(80, 130) / 255, rng.uniform(190, 235) / 255) for _ in range(n_haze)]
    gs += rng.uniform(30, 60, n_haze).tolist()
    for nd in nodes:
        k = int(max(4, min(18, round(px[nd["id"]] * 0.5))))
        sigma = span * (0.035 + px[nd["id"]] * 0.003)
        gx += (nd["x"] + rng.normal(0, sigma, k)).tolist()
        gy += (nd["y"] + rng.normal(0, sigma, k)).tolist()
        gz += (nd["z"] * GALAXY_DISK + rng.normal(0, sigma * 0.5, k)).tolist()
        gc += [tint[nd["id"]]] * k
        gs += rng.uniform(10, 30, k).tolist()
    links = [(by_id[a], by_id[b]) for a, b in g["edges"] if a in by_id and b in by_id]
    for p, q in links:
        length = float(np.hypot(np.hypot(q["x"] - p["x"], q["y"] - p["y"]), (q["z"] - p["z"]) * GALAXY_DISK))
        k = int(max(4, min(24, round(length / span * 40))))
        t = rng.uniform(0.08, 0.92, k)
        gx += (p["x"] + t * (q["x"] - p["x"]) + rng.normal(0, span * 0.025, k)).tolist()
        gy += (p["y"] + t * (q["y"] - p["y"]) + rng.normal(0, span * 0.025, k)).tolist()
        gz += ((p["z"] + t * (q["z"] - p["z"])) * GALAXY_DISK + rng.normal(0, span * 0.0125, k)).tolist()
        cp, cq = np.array(tint[p["id"]]), np.array(tint[q["id"]])
        gc += [tuple(cp + (cq - cp) * ti) for ti in t]
        gs += rng.uniform(8, 20, k).tolist()
    hx, hy = project(gx, gy, gz)
    ax.scatter(hx, hy, s=(np.array(gs) * scale) ** 2, c=gc, alpha=0.05, linewidths=0, zorder=2)

    for p, q in links:
        lx, ly = project([p["x"], q["x"]], [p["y"], q["y"]], [p["z"] * GALAXY_DISK, q["z"] * GALAXY_DISK])
        ax.plot(lx, ly, color="#96B9EB", lw=0.45, alpha=0.55, zorder=3)
    order = sorted(nodes, key=lambda nd: -nd["size"])                 # big groups first, small ones on top
    nx_, ny_ = project([nd["x"] for nd in order], [nd["y"] for nd in order], [nd["z"] * GALAXY_DISK for nd in order])
    cols = [tint[nd["id"]] for nd in order]
    size = np.array([px[nd["id"]] for nd in order])
    ax.scatter(nx_, ny_, s=(size * 2.2 * scale) ** 2, c=cols, alpha=0.12, linewidths=0, zorder=4)   # glow
    ax.scatter(nx_, ny_, s=(size * scale) ** 2, c=cols, alpha=0.95, edgecolors=(0, 0, 0, 0.4), linewidths=0.3, zorder=5)
    for i, nd in enumerate(shown, 1):
        tx, ty = project([nd["x"]], [nd["y"]], [nd["z"] * GALAXY_DISK])
        ax.annotate(f"{'⚠' if gv['risky'](nd) else ''}{i}", (tx[0], ty[0]), xytext=(0, px[nd["id"]] * scale * 0.5 + 1.5),
                    textcoords="offset points", color="#FFF0C8", fontsize=6.5, ha="center", va="bottom", zorder=6)

    half = span * 1.25                                               # frame the disk, stars fill the corners
    ax.set_xlim(-half, half)
    ax.set_ylim(-half * 4.2 / 7.2, half * 4.2 / 7.2)
    ax.set_aspect("equal")
    ax.set_axis_off()
    fig.text(0.5, 0.035, f"Gold = low {gv['color_title']}   ·   Teal = medium   ·   Red = high", color="#878E9C",
             fontsize=7, ha="center")
    fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor=GALAXY_BG)
    plt.close(fig)
    return buf.getvalue()


def _mapper_story(g: dict | None) -> list:
    if not g or not g.get("nodes"):
        return []
    gv = _graph_view(g)
    shown = _labelled(g, gv)
    records = sum(nd["size"] for nd in g["nodes"])
    head = _th(["#", "Group", "What stands out", "Records", *[label for label, _ in gv["columns"]]])
    rows = [head] + [[f"{'⚠ ' if gv['risky'](nd) else ''}{i}", Paragraph(nd.get("label") or f"group {nd['id']}", S["cell"]),
                      Paragraph(nd.get("profile") or "–", S["cell"]), num(nd["size"]), *[f(nd) for _, f in gv["columns"]]]
                     for i, nd in enumerate(shown, 1)]
    n_cols = len(gv["columns"])
    t = Table(rows, colWidths=[1 * cm, 4.6 * cm, 4.4 * cm, 1.6 * cm, *[5.4 * cm / max(1, n_cols)] * n_cols], repeatRows=1)
    style = _table_style()
    for i, nd in enumerate(shown, 1):
        if gv["risky"](nd):
            style.add("TEXTCOLOR", (0, i), (0, i), rl_colors.HexColor(RED))
    t.setStyle(style)
    text = (f"The same graph as CortXplorer's TDA Mapper: {num(len(g['nodes']))} groups of similar records, "
            f"{num(len(g['edges']))} connections, {num(records)} record memberships (a record can sit in overlapping "
            f"groups). Size = records, colour = {gv['color_title']} (teal low → red high). Numbers mark the largest "
            f"groups and the groups at risk ({gv['risk_note']}).")
    galaxy = ("Galaxy view: the same groups and connections flattened into a disk, with gas around the groups and along "
              f"their connections. Colour = {gv['color_title']} (gold low → teal → red high); the numbers are the same "
              "groups as above and in the table.")
    return [KeepTogether(_section("TDA", "TDA Mapper — the shape of the data", text) + [_img(_mapper_png(g, gv, shown))]),
            Spacer(1, 6), KeepTogether([Paragraph(galaxy, S["small"]), Spacer(1, 4), _img(_galaxy_png(g, gv, shown))]),
            Spacer(1, 6), t]


def _mapper_sheet(wb: Workbook, g: dict | None) -> None:
    if not g or not g.get("nodes"):
        return
    gv = _graph_view(g)
    labelled = {nd["id"]: i for i, nd in enumerate(_labelled(g, gv), 1)}
    ws = wb.create_sheet("TDA Mapper groups")
    _head(ws, 1, ["#", "group id", "label", "what stands out", "records", "at risk", gv["color_title"],
                  *[label for label, _ in gv["columns"]], "connections"])
    degree: dict[int, int] = {}
    for a, b in g["edges"]:
        degree[a] = degree.get(a, 0) + 1
        degree[b] = degree.get(b, 0) + 1
    for i, nd in enumerate(sorted(g["nodes"], key=lambda nd: -nd["size"]), 2):
        row = [labelled.get(nd["id"]), nd["id"], nd.get("label") or f"group {nd['id']}", nd.get("profile"), nd["size"],
               "yes" if gv["risky"](nd) else "", round(float(gv["color"](nd)), 4), *[f(nd) for _, f in gv["columns"]],
               degree.get(nd["id"], 0)]
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v)
    _widths(ws, [5, 9, 34, 34, 10, 8, 14] + [14] * len(gv["columns"]) + [12])

    # the two pictures of the graph: Galaxy view and Mapper view (numbers = column "#" of the groups sheet)
    shown = _labelled(g, gv)
    ws = wb.create_sheet("TDA Galaxy & Mapper")
    ws["A1"] = "TDA Galaxy — the Mapper groups as a galaxy"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Colour = {gv['color_title']} (gold low → teal → red high). Size = records. The numbers are the labelled "
                "groups of the sheet 'TDA Mapper groups'.")
    ws["A32"] = "TDA Mapper — the shape of the data"
    ws["A32"].font = Font(bold=True, size=14)
    ws["A33"] = f"Colour = {gv['color_title']} (teal low → red high). Same groups, CortXplorer's Mapper layout."
    for anchor, png in (("A4", _galaxy_png(g, gv, shown)), ("A35", _mapper_png(g, gv, shown))):
        img = XLImage(io.BytesIO(png))
        scale = 820 / img.width
        img.width, img.height = 820, int(img.height * scale)
        ws.add_image(img, anchor)


def _context_sheet(wb: Workbook, ctx: dict | None) -> None:
    if not ctx:
        return
    ws = wb.create_sheet("TDA-ML context")
    for row, (key, value) in enumerate(ctx.items(), 1):
        ws.cell(row=row, column=1, value=key.replace("_", " ")).font = Font(bold=True)
        ws.cell(row=row, column=2, value="; ".join(map(str, value)) if isinstance(value, list) else
                (", ".join(f"{a}={b}" for a, b in value.items()) if isinstance(value, dict) else value))
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
    _widths(ws, [22, 100])


# ── scenario (any dataset) ──────────────────────────────────────────────────
TONES = {"blue": BLUE, "good": GREEN, "warm": ORANGE, "bad": RED}


def _scenario_title(sc: dict) -> str:
    return sc["config"].get("title") or "Scenario simulation"


def _scenario_blocks(sc: dict) -> dict:
    k, c, s = sc["summary"]["kpi"], sc["config"], sc["summary"]
    period, unit, records = c["period_label"], s["dashboard"]["unit"], c["record_label"]
    cap = records[:1].upper() + records[1:]
    inputs = [(f"{cap} in data", f"{num(c['n_records'])} · {num(c['n_periods'])} {unit}" if c["n_periods"] else num(c["n_records"]))]
    outputs = []
    for m in s["metrics"]:
        key, fmt, u = m["key"], m["fmt"], m.get("unit")
        q = s["inputs"]["metrics"].get(key)
        inputs.append((m["label"] + (f" ({u})" if u else ""),
                       f"{_fv(q['p5'], m['record_fmt'])} – {_fv(q['p95'], m['record_fmt'])}" if q else "–"))
        text = (f"expected {m['label']} ({'total' if m['agg'] == 'sum' else 'average'} per {period}); 90% of {unit}: "
                f"{_fv(k.get(f'{key}_p5'), fmt, u)} – {_fv(k.get(f'{key}_p95'), fmt, u)}")
        if m.get("alert") is not None:
            text += f" · P(> {_fv(m['alert'], fmt, u)}) = {pct(k.get(f'{key}_p_over'))}"
        outputs.append(("info", _fv(k.get(f"{key}_mean"), fmt, u), text))
    if s["anomalies"]["high_risk"]:
        outputs.append(("bad" if (k.get("p_any_high") or 0) > 0.5 else "info", pct(k.get("p_any_high")),
                        f"probability a {period} contains at least one high-risk record (anomaly score ≥ {c['high_anomaly']:g})"))
        outputs.append(("info", _adaptive(k.get("high_mean")),
                        f"high-risk {records} expected per {period} (P95 {num(k.get('high_p95'))} of {num(c['period_size'])})"))
    part = lambda p, unit_: p["text"] if "text" in p else _fv(k.get(p["key"]) if p.get("key") else p.get("value"), p["fmt"], unit_)
    tiles = [(t["label"], _fv(k.get(t["key"]), t["fmt"], t.get("unit")),
              "".join(part(p, t.get("unit")) for p in t.get("sub") or []), t.get("tone", "blue"))
             for t in s["dashboard"]["tiles"][:4]]
    how = (f"Each scenario draws a real {period} and resamples its {records}" if c["n_periods"]
           else f"Each scenario resamples {num(c['period_size'])} {records} from the whole dataset")
    return {"inputs": inputs, "inputs_sub": "per record, P5 – P95", "outputs": outputs, "tiles": tiles,
            "engine": (num(sc["n_completed"]), unit,
                       [f"{num(c['n_pool'])} of {num(c['n_records'])} records · {num(c['period_size'])} {records} per {period}", how])}


def _scenario_charts(sc: dict) -> dict[str, Any]:
    s, k = sc["summary"], sc["summary"]["kpi"]
    charts = s["dashboard"]["charts"]
    out: dict[str, Any] = {"dists": []}
    for i in range(0, len(charts), 2):
        pair = charts[i:i + 2]
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.7))
        for ax, c in zip(axes, pair, strict=False):
            lines = []
            for line in c["lines"]:
                v = k.get(line["key"]) if line.get("key") else line.get("value")
                if v is not None:
                    lines.append((v, f"{line['label']} {_fv(v, line['fmt'], c.get('unit'))}", TONES.get(line["tone"], BLUE)))
            _bars(ax, s["hist"][c["key"]], TONES.get(c["tone"], BLUE), lines, c["x_title"])
        for ax in axes[len(pair):]:
            ax.set_axis_off()
        fig.tight_layout()
        out["dists"].append(_png(fig))
    cv = s["convergence"]
    if cv["n"]:
        fig, ax = plt.subplots(figsize=(7.2, 2.2))
        ax.fill_between(cv["n"], cv["lo"], cv["hi"], color=ORANGE, alpha=0.2, label="95% confidence interval")
        ax.plot(cv["n"], cv["p"], color=ORANGE, lw=2, label=cv.get("label", "estimate"))
        ax.set_xscale("log")
        ax.set_xlabel(f"simulated {s['dashboard']['unit']} (log scale)")
        ax.legend(frameon=False, fontsize=8, loc="upper right")
        out["convergence"] = _png(fig)
    return out


def _scenario_story(sc: dict, job: dict) -> list:
    c, s = sc["config"], sc["summary"]
    png = _scenario_charts(sc)
    period, unit, records = c["period_label"], s["dashboard"]["unit"], c["record_label"]
    ctx = s.get("context") or {}
    story: list = [
        Paragraph(f"{_scenario_title(sc).upper()} · MONTE CARLO ON TDA / ML RESULTS", S["eyebrow"]),
        Paragraph("Monte Carlo Simulation with TDA", S["title"]),
        Spacer(1, 4),
        Paragraph(f"Dataset {job['dataset_id']} · job {job['job_id']} · {num(sc['n_completed'])} simulated {unit} · "
                  f"generated {datetime.now(UTC):%d %b %Y}", S["small"]),
    ]
    blocks = _scenario_blocks(sc)
    how = (f"draws one of the {num(c['n_periods'])} historical {unit} and resamples {num(c['period_size'])} of its "
           f"{records}" if c["n_periods"] else f"resamples {num(c['period_size'])} {records} from the whole dataset")
    story += [Spacer(1, 10), _flow(blocks), Spacer(1, 8), _tiles(blocks), Spacer(1, 8),
              Paragraph(f"Each simulated {period} {how} ({num(c['n_pool'])} of {num(c['n_records'])} records in the "
                        "pool). Nothing is fitted: the distributions are the data's own.", S["small"])]
    if png["dists"]:
        story += [KeepTogether(_section("Charts", "Distributions") + [_img(png["dists"][0])])]
        story += [_img(p) for p in png["dists"][1:]]

    shown = s["metrics"][:2]
    tail = s["tails"][0] if s.get("tails") else None
    head = ["Regime (TDA cluster)", "Share", "Avg anomaly", "High-risk", *[m["label"] for m in shown]]
    if tail:
        head.append(f"In high-{tail['label']} {unit}")
    reg = [_th(head)] + [
        [Paragraph(g["label"], S["cell"]), pct(g["share"]), _adaptive(g.get("anomaly_mean")), pct(g.get("high_risk_share")),
         *[_fv(g["metrics"].get(m["key"]), m["record_fmt"]) for m in shown],
         *([f"×{g['tail_lift'][tail['key']]:.2f}" if g["tail_lift"].get(tail["key"]) is not None else "–"] if tail else [])]
        for g in sorted(sc["results"], key=lambda g: -(g.get("high_risk_share") or 0) * g["share"])[:MAX_REGIME_ROWS]]
    extra = len(head) - 4
    tr = Table(reg, colWidths=[4.6 * cm, 1.6 * cm, 1.8 * cm, 1.8 * cm, *[7.2 * cm / max(1, extra)] * extra], repeatRows=1)
    tr.setStyle(_table_style())
    note = (f"The last column shows how over-represented a regime is in the {unit} whose {tail['label']} is in the top 5% "
            "(×1 = as often as usual)." if tail else "")
    note = " ".join(x for x in (note, _regime_note(len(sc["results"]), "by their number of high-risk records")) if x) or None
    story += [KeepTogether(_section("TDA", "Where the risk comes from — TDA regimes", note) + [tr])]
    story += _mapper_story(job.get("graph"))

    notes = []
    if ctx:
        notes.append(f"TDA / ML: {num(ctx.get('n_clusters'))} clusters, {num(ctx.get('n_noise'))} noise records, "
                     f"{num(ctx.get('n_anomalies_high'))} records scored HIGH by the anomaly model.")
        for title, items in (("Relationships", ctx.get("relationships")), ("Patterns", ctx.get("patterns"))):
            if items:
                notes.append(f"{title}: " + "; ".join(items))
    a = s["anomalies"]
    notes.append("Anomaly exclusion off: all records used." if a["threshold"] is None else
                 f"{num(a['excluded'])} records with anomaly score ≥ {a['threshold']} left out of the sampling pool.")
    story += [KeepTogether(_section("Context", "TDA / ML findings and data quality")
                           + [Paragraph(n, S["body"]) for n in notes])]
    if "convergence" in png:
        story += [KeepTogether(_section("Quality", "Are the scenarios enough?") + [_img(png["convergence"])])]
    return story


def _scenario_xlsx(sc: dict, job: dict, wb: Workbook) -> None:
    k, c, s = sc["summary"]["kpi"], sc["config"], sc["summary"]
    unit = s["dashboard"]["unit"]
    xfmt = lambda fmt: "0.0%" if fmt == "pct" else "#,##0" if fmt == "int" else "#,##0.00"
    title = f"Monte Carlo Simulation with TDA — {_scenario_title(sc)}"
    subtitle = (f"Dataset {job['dataset_id']} · job {job['job_id']} · {sc['n_completed']:,} simulated {unit} · "
                f"{c['period_size']:,} {c['record_label']} per {c['period_label']} · {c['method']}")
    _report_sheet(wb, title, subtitle, _scenario_blocks(sc))
    ws = wb.create_sheet("Summary")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = subtitle
    _head(ws, 4, [f"Metric (per {c['period_label']})", "Aggregation", "Unit", "Expected", "P5", "Median", "P95", "Alert",
                  "P(> alert)"])
    row = 5
    for m in s["metrics"]:
        key = m["key"]
        ws.cell(row=row, column=1, value=m["label"])
        ws.cell(row=row, column=2, value="total" if m["agg"] == "sum" else "average")
        ws.cell(row=row, column=3, value=m.get("unit"))
        for j, part in enumerate(("mean", "p5", "p50", "p95", "alert"), 4):
            ws.cell(row=row, column=j, value=k.get(f"{key}_{part}")).number_format = xfmt(m["fmt"])
        ws.cell(row=row, column=9, value=k.get(f"{key}_p_over")).number_format = "0.0%"
        row += 1
    for name, value, fmt in ((f"High-risk records per {c['period_label']} (mean)", k.get("high_mean"), "#,##0.00"),
                             (f"High-risk records per {c['period_label']} (P95)", k.get("high_p95"), "#,##0"),
                             ("High-risk share", k.get("high_share_mean"), "0.0%"),
                             (f"P(a {c['period_label']} has a high-risk record)", k.get("p_any_high"), "0.0%")):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=4, value=value).number_format = fmt
        row += 1
    _widths(ws, [44, 13, 8, 14, 14, 14, 14, 14, 12])

    ws = wb.create_sheet("Inputs (observed)")
    _head(ws, 1, ["Input (per record)", "P5", "Median", "P95"])
    for i, m in enumerate(s["metrics"], 2):
        ws.cell(row=i, column=1, value=m["label"])
        for j, part in enumerate(("p5", "median", "p95"), 2):
            ws.cell(row=i, column=j, value=(s["inputs"]["metrics"].get(m["key"]) or {}).get(part)).number_format = \
                xfmt(m["record_fmt"])
    _widths(ws, [34, 14, 14, 14])

    ws = wb.create_sheet("TDA regimes")
    tails = s.get("tails") or []
    _head(ws, 1, ["regime", "label", "records", "share", "avg anomaly", "high-risk share",
                  *[m["label"] for m in s["metrics"]], *[f"lift in high-{t['label']} {unit}" for t in tails]])
    for i, g in enumerate(sc["results"], 2):
        base = [g["regime"], g["label"], g["records"], g["share"], g.get("anomaly_mean"), g.get("high_risk_share")]
        for j, v in enumerate(base, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if j in (4, 6):
                cell.number_format = "0.0%"
            elif j == 5:
                cell.number_format = "0.000"
        col = len(base) + 1
        for m in s["metrics"]:
            ws.cell(row=i, column=col, value=g["metrics"].get(m["key"])).number_format = xfmt(m["record_fmt"])
            col += 1
        for t in tails:
            ws.cell(row=i, column=col, value=g["tail_lift"].get(t["key"])).number_format = "0.00"
            col += 1
    _widths(ws, [8, 36, 10, 10, 12, 14] + [16] * (len(s["metrics"]) + len(tails)))

    _mapper_sheet(wb, job.get("graph"))

    ws = wb.create_sheet("Distributions")
    col = 1
    for ch in s["dashboard"]["charts"]:
        h = s["hist"][ch["key"]]
        total = sum(h["counts"]) or 1
        ws.cell(row=1, column=col, value=ch["title"]).font = Font(bold=True)
        _head(ws, 2, ["From", "To", "Share"], col)
        for i, (a, b, cnt) in enumerate(zip(h["edges"], h["edges"][1:], h["counts"], strict=False), 3):
            ws.cell(row=i, column=col, value=a)
            ws.cell(row=i, column=col + 1, value=b)
            ws.cell(row=i, column=col + 2, value=cnt / total).number_format = "0.0%"
        chart = BarChart()
        chart.title, chart.legend, chart.gapWidth = ch["title"], None, 20
        chart.add_data(Reference(ws, min_col=col + 2, min_row=2, max_row=2 + len(h["counts"])), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=col, min_row=3, max_row=2 + len(h["counts"])))
        chart.height, chart.width = 7, 13
        ws.add_chart(chart, f"{get_column_letter(col)}{len(h['counts']) + 5}")
        col += 4

    if s["convergence"]["n"]:
        _convergence_sheet(wb, s["convergence"], s["convergence"].get("label", "estimate"), "#,##0.00")
    _context_sheet(wb, s.get("context"))


def _scenario_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("DejaVu", 7)
    canvas.setFillColor(rl_colors.HexColor("#5f6b66"))
    canvas.drawString(2 * cm, 1.2 * cm, "Monte Carlo Simulation with TDA")
    canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
    canvas.restoreState()


# ── loops ───────────────────────────────────────────────────────────────────
def _loops_story(lp: dict, job: dict) -> list:
    s, c = lp["summary"], lp["config"]
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    h = s["null_hist"]
    if h["counts"]:
        _bars(ax, h, ORANGE, [(s["top_persistence"], f"observed {s['top_persistence']:.3f}", BLUE),
                              (s["noise_band"], f"{pct(1 - c['alpha'], 0)} band {s['noise_band']:.3f}", "#7c3aed")],
              "longest H1 persistence in random data")
    png = _png(fig)
    rows = [["#", "Birth", "Death", "Persistence", "p-value", "Verdict"]] + [
        [r["index"] + 1, f"{r['birth']:.4f}", f"{r['death']:.4f}", f"{r['persistence']:.4f}",
         f"{r['p_value']:.4f}", "significant" if r["significant"] else "noise"] for r in lp["results"][:20]]
    t = Table(rows, colWidths=[1.2 * cm, 3 * cm, 3 * cm, 3.2 * cm, 3 * cm, 3.2 * cm], repeatRows=1)
    t.setStyle(_table_style())
    return [
        Paragraph("TOPOLOGY · MONTE CARLO SIGNIFICANCE", S["eyebrow"]),
        Paragraph("Are the loops real?", S["title"]),
        Paragraph(f"Dataset {job['dataset_id']} · job {job['job_id']} · {num(lp['n_completed'])} simulations · "
                  f"null model {c['null']} · {num(c['sample_n'])} sampled points · generated {datetime.now(UTC):%d %b %Y}",
                  S["small"]),
        *_section("Result", f"{s['n_significant']} of {s['n_loops_observed']} loops are significant (p ≤ {c['alpha']})",
                  "Each loop is compared with the longest loop found in random data of the same shape; a loop is "
                  "significant when random data rarely produces one that persists as long."),
        _img(png), Spacer(1, 6), t,
    ]


def _loops_xlsx(lp: dict, job: dict, wb: Workbook) -> None:
    s, c = lp["summary"], lp["config"]
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Monte Carlo significance of H1 loops"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Dataset {job['dataset_id']} · job {job['job_id']}"
    for i, (name, value) in enumerate([("Simulations", lp["n_completed"]), ("Null model", c["null"]),
                                       ("Sample size", c["sample_n"]), ("Alpha", c["alpha"]),
                                       ("Loops observed", s["n_loops_observed"]), ("Significant loops", s["n_significant"]),
                                       ("Longest loop persistence", s["top_persistence"]),
                                       ("Noise band", s["noise_band"])], 4):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=value)
    _widths(ws, [30, 18])
    ws = wb.create_sheet("Loops")
    _head(ws, 1, ["#", "Birth", "Death", "Persistence", "p-value", "Significant"])
    for i, r in enumerate(lp["results"], 2):
        for j, v in enumerate([r["index"] + 1, r["birth"], r["death"], r["persistence"], r["p_value"], r["significant"]], 1):
            ws.cell(row=i, column=j, value=v)
    _widths(ws, [6, 12, 12, 14, 12, 12])


# ── shared ──────────────────────────────────────────────────────────────────
def _head(ws, row: int, labels: list[str], col: int = 1) -> None:
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=col + i, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _convergence_sheet(wb: Workbook, cv: dict, label: str, number_format: str = "0.0%") -> None:
    ws = wb.create_sheet("Convergence")
    _head(ws, 1, ["Simulations", label, "95% CI lower", "95% CI upper"])
    for i, row in enumerate(zip(cv["n"], cv["p"], cv["lo"], cv["hi"], strict=False), 2):
        ws.cell(row=i, column=1, value=row[0])
        for j, v in enumerate(row[1:], 2):
            ws.cell(row=i, column=j, value=v).number_format = number_format
    chart = LineChart()
    chart.title = "Estimate as simulations grow"
    chart.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=1 + len(cv["n"])), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=1 + len(cv["n"])))
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "F2")
    _widths(ws, [12, 18, 14, 14])


def kind(result: dict[str, Any]) -> str:
    for name in ("fleet", "scenario", "loops"):
        if result.get(name):
            return name
    raise ValueError("This job has no fleet, scenario or loop result to report on.")


def build_pdf(result: dict[str, Any], graph: dict | None = None) -> bytes:
    """``graph``: the TDA Mapper graph the job was sent (from its live feed), if any."""
    job = {"dataset_id": result["dataset_id"], "job_id": result["job_id"], "graph": graph}
    what = kind(result)
    story = {"fleet": _fleet_story, "scenario": _scenario_story, "loops": _loops_story}[what](result[what], job)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                            bottomMargin=2 * cm, title="Monte Carlo report", author="CortXplorer Monte Carlo Service")
    footer = _scenario_footer if what == "scenario" else _footer
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_xlsx(result: dict[str, Any], graph: dict | None = None) -> bytes:
    job = {"dataset_id": result["dataset_id"], "job_id": result["job_id"], "graph": graph}
    wb = Workbook()
    what = kind(result)
    {"fleet": _fleet_xlsx, "scenario": _scenario_xlsx, "loops": _loops_xlsx}[what](result[what], job, wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
