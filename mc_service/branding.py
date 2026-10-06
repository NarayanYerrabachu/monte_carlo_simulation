"""Monte Carlo logo and name for every file the service produces.

* ``brand_page(callback)``  wraps a ReportLab page callback: draws the mark and
  the name at the top of every page, inside the top margin, then runs the callback.
* ``brand_workbook(wb)``    puts the logo at the top of every sheet of a finished
  workbook: each sheet moves down by two rows, and what points at cells (merged
  ranges, frozen panes, filters, conditional formats, chart and image anchors)
  is moved with it.
* ``logo_png(kind, height)`` PNG of the mark or of the mark + name.

The mark's geometry mirrors ``frontend/img/montecarlo-mark.svg`` (the web logo):
one starting point fanning out into five simulated paths; the amber one is the
typical outcome. If the SVG changes, change the numbers here too.
"""
from __future__ import annotations

import functools
import io
from typing import Any, Callable

NAME = "Monte Carlo"
TAGLINE = "SIMULATED OUTCOMES"
TEAL, PURPLE, AMBER = "#5E9CA6", "#9B7FD4", "#E0A33E"
NAVY, PRINT_PURPLE, GREY = "#1B2A4A", "#6B4FBB", "#6B7280"

# Geometry on a 64 × 64 canvas (y down), as in the SVG: (colour, path, end-node radius).
_START = (10, 32)
_PATHS = [
    (PURPLE, [(10, 32), (24, 26), (38, 17), (54, 10)], 3.2),
    (PURPLE, [(10, 32), (25, 39), (39, 48), (54, 54)], 3.2),
    (TEAL, [(10, 32), (25, 33), (40, 24), (54, 21)], 3.9),
    (TEAL, [(10, 32), (24, 36), (40, 40), (54, 43)], 3.9),
    (AMBER, [(10, 32), (26, 29), (40, 34), (54, 32)], 5.0),
]
_STROKE, _R_START = 2.6, 4.2


# ── PDF (ReportLab canvas, vector) ───────────────────────────────────────────

def draw_mark(canvas, x: float, y: float, size: float) -> None:
    """Draw the mark with its lower-left corner at (x, y), ``size`` points wide and high."""
    from reportlab.lib.colors import HexColor
    k = size / 64.0
    pt = lambda p: (x + p[0] * k, y + (64 - p[1]) * k)            # noqa: E731  (PDF y axis points up)
    canvas.saveState()
    canvas.setLineCap(1); canvas.setLineJoin(1); canvas.setLineWidth(_STROKE * k)
    for colour, path, r in _PATHS:
        canvas.setStrokeColor(HexColor(colour))
        p = canvas.beginPath()
        p.moveTo(*pt(path[0]))
        for q in path[1:]:
            p.lineTo(*pt(q))
        canvas.drawPath(p, stroke=1, fill=0)
        canvas.setFillColor(HexColor(colour))
        canvas.circle(*pt(path[-1]), r * k, stroke=0, fill=1)
    canvas.setFillColor(HexColor(TEAL))
    canvas.circle(*pt(_START), _R_START * k, stroke=0, fill=1)
    canvas.restoreState()


def draw_page_header(canvas, doc) -> None:
    """Mark, name and tagline at the top of a page, with a thin rule underneath."""
    from reportlab.lib.colors import HexColor
    from reportlab.lib.units import mm
    width, height = doc.pagesize
    left, right = doc.leftMargin, width - doc.rightMargin
    size = 6.8 * mm
    base = height - 11.0 * mm                                   # header band: 4.2 – 11.0 mm from the top edge
    canvas.saveState()
    draw_mark(canvas, left, base, size)
    x, y = left + size + 2.4 * mm, base + 2.0 * mm
    for text, colour in (("Monte ", NAVY), ("C", PRINT_PURPLE), ("arlo", NAVY)):
        canvas.setFont("Helvetica-Bold", 12.5); canvas.setFillColor(HexColor(colour))
        canvas.drawString(x, y, text)
        x += canvas.stringWidth(text, "Helvetica-Bold", 12.5)
    canvas.setFont("Helvetica", 6.3); canvas.setFillColor(HexColor(GREY))
    canvas.drawRightString(right, y + 0.3 * mm, " ".join(TAGLINE).replace("   ", "     "))
    canvas.setStrokeColor(HexColor("#D9D4EA")); canvas.setLineWidth(0.5)
    canvas.line(left, base - 1.0 * mm, right, base - 1.0 * mm)
    canvas.restoreState()


def brand_page(callback: Callable[[Any, Any], None] | None = None) -> Callable[[Any, Any], None]:
    """A ReportLab page callback that draws the brand header, then runs ``callback``."""
    def page(canvas, doc):
        draw_page_header(canvas, doc)
        if callback is not None:
            callback(canvas, doc)
    return page


# ── PNG (for Excel) ──────────────────────────────────────────────────────────

def _font(size: int, bold: bool):
    from PIL import ImageFont
    try:                                                         # DejaVu ships with matplotlib: present in every deployment
        from matplotlib import font_manager
        return ImageFont.truetype(font_manager.findfont("DejaVu Sans:bold" if bold else "DejaVu Sans"), size)
    except Exception:
        return ImageFont.load_default()


@functools.lru_cache(maxsize=8)
def logo_png(kind: str = "lockup", height: int = 44) -> bytes:
    """PNG of the logo on a transparent background. kind: 'mark' or 'lockup' (mark + name + tagline)."""
    from PIL import Image, ImageDraw
    ss = 4                                                       # draw large, scale down: smooth edges
    h = height * ss
    k = h / 64.0
    wordmark = kind == "lockup"
    name_font, tag_font = _font(int(h * 0.50), True), _font(int(h * 0.17), False)
    probe = ImageDraw.Draw(Image.new("RGBA", (10, 10)))
    text_w = int(probe.textlength(NAME, font=name_font)) if wordmark else 0
    w = h + (int(h * 0.30) + text_w + int(h * 0.10) if wordmark else 0)
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    pt = lambda p: (p[0] * k, p[1] * k)                           # noqa: E731
    lw = max(1, int(round(_STROKE * k)))
    for colour, path, r in _PATHS:
        d.line([pt(p) for p in path], fill=colour, width=lw, joint="curve")
        cx, cy = pt(path[-1])
        d.ellipse([cx - r * k, cy - r * k, cx + r * k, cy + r * k], fill=colour)
    cx, cy = pt(_START)
    d.ellipse([cx - _R_START * k, cy - _R_START * k, cx + _R_START * k, cy + _R_START * k], fill=TEAL)
    if wordmark:
        x, y = h + int(h * 0.30), int(h * 0.10)
        for text, colour in (("Monte ", NAVY), ("C", PRINT_PURPLE), ("arlo", NAVY)):
            d.text((x, y), text, font=name_font, fill=colour)
            x += int(d.textlength(text, font=name_font))
        d.text((h + int(h * 0.32), int(h * 0.72)), " ".join(TAGLINE), font=tag_font, fill=GREY)
    out = img.resize((max(1, w // ss), height), Image.LANCZOS)
    buf = io.BytesIO()
    out.save(buf, format="PNG")
    return buf.getvalue()


# ── Excel (openpyxl) ─────────────────────────────────────────────────────────

HEADER_ROWS = 2        # rows the logo takes at the top of every sheet


def _shift_ref(ref: str, n: int) -> str:
    from openpyxl.utils import get_column_letter, range_boundaries
    out = []
    for part in str(ref).split():
        c1, r1, c2, r2 = range_boundaries(part)
        a, b = f"{get_column_letter(c1)}{r1 + n}", f"{get_column_letter(c2)}{r2 + n}"
        out.append(a if a == b and ":" not in part else f"{a}:{b}")
    return " ".join(out)


def _shift_anchor(obj, n: int) -> None:
    """Move a chart's or image's anchor down by ``n`` rows (a cell string, or an anchor object)."""
    anchor = getattr(obj, "anchor", None)
    if isinstance(anchor, str):
        obj.anchor = _shift_ref(anchor, n)
    elif anchor is not None and getattr(anchor, "_from", None) is not None:
        anchor._from.row += n
        if getattr(anchor, "to", None) is not None:
            anchor.to.row += n


def brand_sheet(ws) -> None:
    """Logo at the top of a finished sheet: everything moves down by HEADER_ROWS rows."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.utils.cell import coordinate_from_string

    n = HEADER_ROWS
    merged = [str(r) for r in ws.merged_cells.ranges]
    for r in merged:
        ws.unmerge_cells(r)
    freeze, filt, rules = ws.freeze_panes, ws.auto_filter.ref, list(ws.conditional_formatting)
    heights = {row: dim.height for row, dim in ws.row_dimensions.items() if dim.height}
    ws.insert_rows(1, n)
    for r in merged:
        ws.merge_cells(_shift_ref(r, n))
    if freeze:
        col, row = coordinate_from_string(freeze)
        ws.freeze_panes = f"{col}{row + n}"
    if filt:
        ws.auto_filter.ref = _shift_ref(filt, n)
    if rules:
        moved = ConditionalFormattingList()
        for cf in rules:
            for rule in cf.rules:
                moved.add(_shift_ref(cf.sqref, n), rule)
        ws.conditional_formatting = moved
    for row, height in heights.items():
        ws.row_dimensions[row].height = None
    for row, height in heights.items():
        ws.row_dimensions[row + n].height = height
    for obj in list(getattr(ws, "_charts", [])) + list(getattr(ws, "_images", [])):
        _shift_anchor(obj, n)
    try:
        ws.add_image(XLImage(io.BytesIO(logo_png("lockup", 44))), "A1")
        ws.row_dimensions[1].height = 36                         # points; the 44 px logo fits
    except Exception:                                            # no image support → the name as text
        ws.cell(row=1, column=1, value=NAME)


def brand_workbook(wb, title: str = "") -> None:
    """Logo on every sheet, name in every printed page header, page numbers in the footer."""
    for ws in wb.worksheets:
        brand_sheet(ws)
        try:
            ws.oddHeader.left.text = f"&B{NAME}&B  ·  Simulated outcomes"
            ws.oddHeader.left.size = 9
            if title:
                ws.oddHeader.right.text = title[:90].replace("&", "&&")
                ws.oddHeader.right.size = 8
            ws.oddFooter.right.text = "Page &P of &N"
            ws.oddFooter.right.size = 8
        except Exception:
            pass
