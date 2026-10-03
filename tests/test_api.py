import base64
import gzip
import io
import json
import threading
import time

import numpy as np
import pytest
from fastapi.testclient import TestClient

from mc_service import main
from mc_service.engine import simulate
from mc_service.simulations import RUNNERS


def data(r):
    body = r.json()
    assert set(body) == {"status", "data", "message"}
    return body["data"]


def _loops_request(**overrides):
    req = {
        "contract_version": "1",
        "dataset_id": "ds-1",
        "settings": {"n_sims": 20, "seed": 1},
        "loops": {"X": [[0.0, 1.0], [1.0, 0.0], [0.5, 0.5]], "sample_n": 50},
    }
    req.update(overrides)
    return req


def _wait(client, job_id, timeout=10.0):
    end = time.time() + timeout
    while time.time() < end:
        st = data(client.get(f"/v1/jobs/{job_id}"))
        if st["status"] not in ("queued", "running"):
            return st
        time.sleep(0.02)
    raise AssertionError("job did not finish")


@pytest.fixture
def client():
    with TestClient(main.app) as c:
        yield c


@pytest.fixture
def fake_loops(monkeypatch):
    """Register a trivial loops runner that uses the real engine."""
    def run(section, ctx):
        res = simulate(lambda rng: rng.normal(), ctx.sim_config())
        return {"test": "loops", "n_completed": res.n_completed, "stopped_early": res.stopped_early,
                "elapsed_s": res.elapsed_s, "config": {"n_points": len(section.X)},
                "results": [], "summary": {"null_mean": float(res.null.mean())}}
    monkeypatch.setitem(RUNNERS, "loops", run)


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200
    assert data(r) == {"status": "ok", "version": main.__version__, "contract_version": "1"}


def test_submit_poll_result(client, fake_loops):
    r = client.post("/v1/jobs", json=_loops_request())
    assert r.status_code == 202
    job_id = data(r)["job_id"]
    st = _wait(client, job_id)
    assert st["status"] == "done"
    assert st["progress"]["loops"] == {"done": 20, "total": 20}
    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["dataset_id"] == "ds-1" and res["job_id"] == job_id
    assert res["loops"]["n_completed"] == 20
    assert res["relationships"] is None and res["errors"] == {}


def test_same_request_same_result(client, fake_loops):
    out = []
    for _ in range(2):
        job_id = data(client.post("/v1/jobs", json=_loops_request()))["job_id"]
        _wait(client, job_id)
        res = data(client.get(f"/v1/jobs/{job_id}/result"))
        out.append(res["loops"]["summary"])
    assert out[0] == out[1]


def test_unimplemented_test_reported_per_test(client, fake_loops):
    req = _loops_request(relationships={"labels_a": ["a", "b"], "labels_b": ["x", "y"],
                                        "node_members": [[0, 1]]})
    job_id = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["loops"] is not None
    assert res["errors"] == {"relationships": "not implemented yet"}


def test_all_tests_failing_marks_job_failed(client, monkeypatch):
    def boom(section, ctx):
        raise RuntimeError("broken")
    monkeypatch.setitem(RUNNERS, "loops", boom)
    job_id = data(client.post("/v1/jobs", json=_loops_request()))["job_id"]
    assert _wait(client, job_id)["status"] == "failed"
    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["errors"] == {"loops": "RuntimeError: broken"}


def test_result_409_while_running_then_cancel(client, monkeypatch):
    started, release = threading.Event(), threading.Event()

    def blocking(section, ctx):
        started.set()
        release.wait(5)
        return simulate(lambda rng: 0.0, ctx.sim_config())   # sees the cancel flag → Cancelled
    monkeypatch.setitem(RUNNERS, "loops", blocking)

    job_id = data(client.post("/v1/jobs", json=_loops_request()))["job_id"]
    assert started.wait(5)
    r = client.get(f"/v1/jobs/{job_id}/result")
    assert r.status_code == 409 and r.json()["status"] == "error"

    assert client.delete(f"/v1/jobs/{job_id}").status_code == 200
    release.set()
    assert _wait(client, job_id)["status"] == "cancelled"
    assert client.get(f"/v1/jobs/{job_id}/result").status_code == 409


def test_unknown_job_404(client):
    r = client.get("/v1/jobs/nope")
    assert r.status_code == 404
    assert r.json() == {"status": "error", "data": None, "message": "Job nope not found (unknown or expired)"}


@pytest.mark.parametrize("change, fragment", [
    ({"contract_version": "2"}, "contract_version"),
    ({"loops": None}, "select at least one test"),
    ({"relationships": {"labels_a": ["a"], "labels_b": ["x", "y"], "node_members": [[0]]}}, "labels_b"),
    ({"relationships": {"labels_a": ["a", "b"], "labels_b": ["x", "y"], "node_members": [[0, 5]]}}, "outside"),
    ({"settings": {"n_sims": 5}}, "n_sims"),
    ({"loops": {"X": [[0.0, 1.0], [1.0]]}}, "different lengths"),
])
def test_invalid_requests_422(client, change, fragment):
    r = client.post("/v1/jobs", json=_loops_request(**change))
    assert r.status_code == 422
    assert fragment in r.json()["message"]


def test_gzip_body_accepted(client, fake_loops):
    raw = gzip.compress(json.dumps(_loops_request()).encode())
    r = client.post("/v1/jobs", content=raw,
                    headers={"Content-Encoding": "gzip", "Content-Type": "application/json"})
    assert r.status_code == 202
    assert _wait(client, data(r)["job_id"])["status"] == "done"


def test_bad_gzip_400(client):
    r = client.post("/v1/jobs", content=b"not gzip",
                    headers={"Content-Encoding": "gzip", "Content-Type": "application/json"})
    assert r.status_code == 400


def test_body_limit_413(client, monkeypatch):
    mw = next(m for m in main.app.user_middleware if m.cls.__name__ == "BodyLimitMiddleware")
    monkeypatch.setitem(mw.kwargs, "max_bytes", 200)
    main.app.middleware_stack = None                     # rebuild with the smaller limit
    try:
        with TestClient(main.app) as small:
            big = _loops_request(loops={"X": [[0.0, 1.0]] * 100})
            assert small.post("/v1/jobs", json=big).status_code == 413
            bomb = gzip.compress(json.dumps(big).encode())   # small compressed, large inflated
            assert len(bomb) < 200
            r = small.post("/v1/jobs", content=bomb,
                           headers={"Content-Encoding": "gzip", "Content-Type": "application/json"})
            assert r.status_code == 413
    finally:
        monkeypatch.undo()
        main.app.middleware_stack = None


def test_real_loops_runner_and_live_endpoint(client):
    rng = np.random.default_rng(0)
    t = rng.uniform(0, 2 * np.pi, 300)
    X = np.column_stack([np.cos(t), np.sin(t), rng.normal(scale=0.05, size=300)]).tolist()
    req = _loops_request(loops={"X": X, "sample_n": 150}, settings={"n_sims": 19, "seed": 3})
    job_id = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, job_id)["status"] == "done"

    live = data(client.get(f"/v1/jobs/{job_id}/live"))
    feed = live["live"]["loops"]
    assert live["status"] == "done" and feed["n_null"] == 19
    assert feed["observed"]["points"] and feed["frame"]["index"] == 18
    assert data(client.get(f"/v1/jobs/{job_id}/live?since=19"))["live"]["loops"]["null"] == []

    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["loops"]["summary"]["n_significant"] == 1
    assert res["loops"]["results"][0]["p_value"] == pytest.approx(feed["running"]["p_value"])


def test_list_jobs_newest_first(client, fake_loops):
    ids = [data(client.post("/v1/jobs", json=_loops_request(dataset_id=f"d{i}")))["job_id"] for i in range(2)]
    listed = [j["job_id"] for j in data(client.get("/v1/jobs"))]
    assert listed.index(ids[1]) < listed.index(ids[0])


def test_viewer_page_served(client):
    r = client.get("/")
    assert r.status_code == 200 and "text/html" in r.headers["content-type"]
    assert "viewer.js" in r.text
    js = client.get("/js/viewer.js")
    assert js.status_code == 200 and js.headers["cache-control"] == "no-cache"


def test_fleet_job_end_to_end(client):
    from tests.test_fleet_sim import _records
    req = {"contract_version": "1", "dataset_id": "fleet-ds", "settings": {"n_sims": 100, "seed": 1},
           "fleet": _records()}
    job_id = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["fleet"]["summary"]["kpi"]["p_meet_sla"] > 0.9 and res["errors"] == {}


def test_fleet_rerun_with_new_parameters(client):
    from tests.test_fleet_sim import _records
    req = {"contract_version": "1", "dataset_id": "fleet-ds", "settings": {"n_sims": 100, "seed": 1},
           "fleet": {**_records(), "context": {"n_clusters": 3}}}
    first = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, first)["status"] == "done"

    r = client.post(f"/v1/jobs/{first}/rerun", json={"n_sims": 60, "fleet": {"fleet_size": 40, "sla_on_time": 0.9}})
    assert r.status_code == 202
    second = data(r)
    assert second["rerun_of"] == first
    assert _wait(client, second["job_id"])["status"] == "done"
    res = data(client.get(f"/v1/jobs/{second['job_id']}/result"))["fleet"]
    assert res["config"]["fleet_size"] == 40 and res["config"]["sla_on_time"] == 0.9
    assert res["n_completed"] == 60 and res["summary"]["context"] == {"n_clusters": 3}
    assert res["summary"]["inputs"]["vehicles"] == 20


def test_rerun_rejects_non_fleet_jobs(client, fake_loops):
    job_id = data(client.post("/v1/jobs", json=_loops_request()))["job_id"]
    _wait(client, job_id)
    r = client.post(f"/v1/jobs/{job_id}/rerun", json={"fleet": {"fleet_size": 5}})
    assert r.status_code == 409 and "only fleet and scenario jobs" in r.json()["message"]
    assert client.post(f"/v1/jobs/{job_id}/rerun", json={"fleet": {"colour": 1}}).status_code == 422


def _scenario_request(**section):
    from tests.test_scenario_sim import _batches
    rows = _batches(bad_period=3)
    rows["mapper"] = {"nodes": [{"id": i, "x": i * 0.3, "y": (i % 3) * 0.4, "z": i * 0.1,
                                 "members": list(range(i * 20, i * 20 + 20)), "label": f"Week {i + 1}"} for i in range(12)],
                      "edges": [[i, i + 1] for i in range(11)]}
    return {"contract_version": "1", "dataset_id": "pharma-ds", "settings": {"n_sims": 100, "seed": 1},
            "scenario": {**rows, "title": "Pharma batches", "period_label": "week", "record_label": "batches", **section}}


def test_scenario_job_end_to_end_with_reports_and_rerun(client):
    """Any table: the job describes its own dashboard; the reports carry the TDA Mapper graph."""
    job_id = data(client.post("/v1/jobs", json=_scenario_request()))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    res = data(client.get(f"/v1/jobs/{job_id}/result"))
    assert res["errors"] == {} and res["scenario"]["config"]["title"] == "Pharma batches"
    live = data(client.get(f"/v1/jobs/{job_id}/live"))["live"]["scenario"]
    assert live["observed"]["dashboard"]["unit"] == "weeks" and live["n_null"] == 100

    rep = data(client.post("/v1/reports/job", json={"job_id": job_id, "format": "json"}))
    assert rep["kind"] == "scenario" and len(rep["graph"]["nodes"]) == 12 and rep["graph"]["view"]["color_key"] == "risk"
    for fmt, magic in (("pdf", b"%PDF"), ("xlsx", b"PK")):
        f = data(client.post("/v1/reports/job", json={"job_id": job_id, "format": fmt}))
        assert f["filename"] == f"monte-carlo-scenario-{job_id[:8]}.{fmt}"
        assert base64.b64decode(f["content"]).startswith(magic)

    r = client.post(f"/v1/jobs/{job_id}/rerun", json={"n_sims": 60, "scenario": {"period_size": 10, "alerts": {"cost": 9000}}})
    assert r.status_code == 202
    again = data(r)["job_id"]
    assert _wait(client, again)["status"] == "done"
    res2 = data(client.get(f"/v1/jobs/{again}/result"))["scenario"]
    assert res2["config"]["period_size"] == 10 and res2["n_completed"] == 60
    assert res2["summary"]["kpi"]["cost_alert"] == 9000 and res2["summary"]["kpi"]["cost_p_over"] == 1.0
    assert client.post(f"/v1/jobs/{job_id}/rerun", json={"scenario": {"colour": 1}}).status_code == 422


def test_scenario_request_validation(client):
    bad = _scenario_request()
    bad["scenario"]["metrics"][0]["values"] = [1.0]
    r = client.post("/v1/jobs", json=bad)
    assert r.status_code == 422 and "metrics[yield_pct].values" in r.json()["message"]
    dup = _scenario_request()
    dup["scenario"]["metrics"][1]["key"] = "yield_pct"
    assert "unique" in client.post("/v1/jobs", json=dup).json()["message"]


def test_fleet_report_carries_the_mapper_graph(client):
    from tests.test_fleet_sim import _records
    rows = _records()
    rows["mapper"] = {"nodes": [{"id": 0, "x": 0, "y": 0, "z": 0, "members": list(range(100)), "label": "Depot A"},
                                {"id": 1, "x": 1, "y": 1, "z": 1, "members": list(range(100, 200)), "label": "Depot B"}],
                      "edges": [[0, 1]]}
    req = {"contract_version": "1", "dataset_id": "fleet-ds", "settings": {"n_sims": 100, "seed": 1}, "fleet": rows}
    job_id = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    rep = data(client.post("/v1/reports/job", json={"job_id": job_id, "format": "json"}))
    assert rep["kind"] == "fleet" and [nd["label"] for nd in rep["graph"]["nodes"]] == ["Depot A", "Depot B"]
    xlsx = base64.b64decode(data(client.post("/v1/reports/job", json={"job_id": job_id, "format": "xlsx"}))["content"])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames[0] == "Report" and "TDA Mapper groups" in wb.sheetnames
    assert len(wb["TDA Galaxy & Mapper"]._images) == 2                  # Galaxy view and Mapper view
    assert base64.b64decode(data(client.post("/v1/reports/job", json={"job_id": job_id, "format": "pdf"}))["content"]).startswith(b"%PDF")


@pytest.mark.parametrize("fmt, magic", [("pdf", b"%PDF"), ("xlsx", b"PK")])
def test_job_report_files(client, fmt, magic):
    import base64

    from tests.test_fleet_sim import _records
    req = {"contract_version": "1", "dataset_id": "fleet-ds", "settings": {"n_sims": 100, "seed": 1},
           "fleet": {**_records(slow_day=2), "context": {"n_clusters": 2, "relationships": ["A ↔ B"]}}}
    job_id = data(client.post("/v1/jobs", json=req))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    f = data(client.post("/v1/reports/job", json={"job_id": job_id, "format": fmt}))
    content = base64.b64decode(f["content"])
    assert content.startswith(magic) and f["size_bytes"] == len(content) and f["filename"].endswith(f".{fmt}")
    j = data(client.post("/v1/reports/job", json={"job_id": job_id}))
    assert j["kind"] == "fleet" and j["result"]["fleet"]["summary"]["kpi"]["fleet_size"] == 20


def test_job_report_for_loops_and_errors(client):
    import base64
    rng = np.random.default_rng(0)
    t = rng.uniform(0, 2 * np.pi, 300)
    X = np.column_stack([np.cos(t), np.sin(t), rng.normal(scale=0.05, size=300)]).tolist()
    job_id = data(client.post("/v1/jobs", json=_loops_request(loops={"X": X, "sample_n": 150},
                                                              settings={"n_sims": 19})))["job_id"]
    assert _wait(client, job_id)["status"] == "done"
    pdf = base64.b64decode(data(client.post("/v1/reports/job", json={"job_id": job_id, "format": "pdf"}))["content"])
    assert pdf.startswith(b"%PDF")
    assert client.post("/v1/reports/job", json={"job_id": "nope"}).status_code == 404


def test_finished_jobs_survive_a_restart(tmp_path):
    from mc_service.contract import SimulationRequest
    from mc_service.jobs import JobStore
    from tests.test_fleet_sim import _records

    store = JobStore(workers=1, n_jobs=1, ttl_s=3600, job_dir=str(tmp_path))
    req = SimulationRequest.model_validate({"contract_version": "1", "dataset_id": "fleet-ds",
                                            "settings": {"n_sims": 50, "seed": 1}, "fleet": _records()})
    job = store.submit(req)
    for _ in range(200):
        if job.status == "done":
            break
        time.sleep(0.02)
    assert job.status == "done" and (tmp_path / f"{job.id}.json.gz").exists()
    store.shutdown()

    reloaded = JobStore(workers=1, n_jobs=1, ttl_s=3600, job_dir=str(tmp_path))      # "container restarted"
    again = reloaded.get(job.id)
    assert again is not None and again.status == "done"
    assert again.response()["fleet"]["summary"]["kpi"] == job.response()["fleet"]["summary"]["kpi"]
    assert again.live["fleet"].snapshot()["n_null"] == 50                           # replay still works
    with pytest.raises(ValueError, match="Monte Carlo ↗"):                          # records not kept across restarts
        reloaded.rerun(again, None, {"fleet": {"fleet_size": 5}})
    reloaded.shutdown()

    expired = JobStore(workers=1, n_jobs=1, ttl_s=0.001, job_dir=str(tmp_path))
    assert expired.get(job.id) is None and not (tmp_path / f"{job.id}.json.gz").exists()
    expired.shutdown()
