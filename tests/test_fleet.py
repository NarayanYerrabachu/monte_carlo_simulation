import base64
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from mc_service import main
from mc_service.fleet.model import build_report, simulate

N = 2_000


@pytest.fixture(scope="module")
def client():
    with TestClient(main.app) as c:
        yield c


def test_report_is_deterministic_per_seed():
    a, b = build_report(N, 1).data, build_report.__wrapped__(N, 1).data   # bypass the cache
    assert a["kpi"] == b["kpi"] and a["text"] == b["text"]
    assert build_report(N, 2).data["kpi"] != a["kpi"]


def test_kpis_are_consistent():
    d = build_report(N, 1).data
    k = d["kpi"]
    assert 0 < k["p_day_on_time"] < 1
    assert k["p_day_on_time_ci"][0] < k["p_day_on_time"] < k["p_day_on_time_ci"][1]
    assert k["service_mean"] >= k["p_day_on_time"]            # a partly late day still delivers most
    assert k["v_required_p50"] <= k["v_required_p95"]
    assert k["total_cost_mean"] == pytest.approx(k["fuel_cost_mean"] + k["maint_cost_mean"] + k["overtime_cost_mean"])
    assert sum(d["service_bands"].values()) == pytest.approx(1)
    assert d["service_bands"]["100%"] == pytest.approx(k["p_day_on_time"])
    assert sum(d["breakdowns_pmf"]) == N
    assert len(d["recommendations"]) == 5 and len(d["findings"]) == 6


def test_more_drivers_never_hurt_in_sizing():
    sz = build_report(N, 1).data["sizing"]
    assert all(s >= f - 0.03 for f, s in zip(sz["drivers_fixed"], sz["drivers_scaled"]))


def test_simulate_backlog_zero_with_huge_fleet():
    import numpy as np
    df = simulate(np.random.default_rng(0), 500, vehicles=2_000, driver_pool=2_500)
    assert (df["backlog"] == 0).all()


def test_report_json_endpoint(client):
    r = client.post("/v1/reports/fleet", json={"n": N, "seed": 1})
    assert r.status_code == 200 and r.headers["content-type"] == "application/json"
    d = r.json()["data"]
    assert d["n"] == N and d["seed"] == 1 and d["kpi"] == build_report(N, 1).data["kpi"]


@pytest.mark.parametrize("body, fragment", [
    ({"n": 10}, "n"),
    ({"n": N, "seed": -1}, "seed"),
    ({"n": N, "colour": "red"}, "colour"),          # unknown fields are rejected
])
def test_report_rejects_bad_json(client, body, fragment):
    r = client.post("/v1/reports/fleet", json=body)
    assert r.status_code == 422
    assert r.json()["status"] == "error" and fragment in r.json()["message"]


def test_report_defaults_from_empty_json(client):
    d = client.post("/v1/reports/fleet", json={}).json()["data"]
    assert d["n"] == 10_000 and d["seed"] == 42


def _file(client, kind):
    r = client.post(f"/v1/reports/fleet/{kind}", json={"n": N, "seed": 1})
    assert r.status_code == 200 and r.headers["content-type"] == "application/json"
    f = r.json()["data"]
    content = base64.b64decode(f["content"])
    assert f["encoding"] == "base64" and f["size_bytes"] == len(content)
    assert f["filename"] == f"fleet-monte-carlo-{N}-seed1.{kind}"
    return f, content


def test_pdf_as_json(client):
    f, content = _file(client, "pdf")
    assert f["content_type"] == "application/pdf"
    assert content.startswith(b"%PDF") and len(content) > 50_000


def test_excel_as_json(client):
    f, content = _file(client, "xlsx")
    assert "spreadsheetml" in f["content_type"]
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Inputs", "Distributions", "Drivers", "Fleet sizing", "Convergence", "Scenarios"]
    assert wb["Scenarios"].max_row == N + 1
    k = build_report(N, 1).data["kpi"]
    assert wb["Summary"]["B5"].value == pytest.approx(k["p_day_on_time"])
    assert wb["Drivers"]._charts and wb["Fleet sizing"]._charts


def test_report_page_and_viewer_link(client):
    page = client.get("/report")
    assert page.status_code == 200 and "Download PDF" in page.text and "Download Excel" in page.text
    assert 'href="/report"' in client.get("/").text


def test_report_page_for_a_job(client):
    page = client.get("/report?job=abc")
    assert page.status_code == 200 and "/js/fleet.js" in page.text and 'id="dl-pdf"' in page.text
    assert "fleet_report.js" in client.get("/report").text                  # no job: assumption report


def test_fleet_dashboard_served(client):
    page = client.get("/fleet")
    assert page.status_code == 200 and "/js/fleet.js" in page.text and "/css/fleet.css" in page.text
    assert client.get("/js/fleet.js").status_code == 200
    assert 'href="/fleet"' in client.get("/").text


def test_frontend_assets_served(client):
    page = client.get("/report").text
    assert "/js/fleet_report.js" in page and "/css/fleet_report.css" in page
    for path in ("/js/fleet_report.js", "/css/fleet_report.css", "/js/viewer.js", "/css/viewer.css"):
        r = client.get(path)
        assert r.status_code == 200 and r.headers["cache-control"] == "no-cache", path
